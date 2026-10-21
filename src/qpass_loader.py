"""Load one QPASS workbook into a standardized long DataFrame.

QPASS workbooks come in three layouts (see docs for details):

- Format A (2018-2019): "wide" weekly sheets. Each sheet holds one week; each
  day is a block of four time columns. Header spread over several rows.
- Format B (2024-2025 to Fall 2025-2026): "long" sheets named
  "QPASS Data Entry (Week N)", one row per location x space type x day.
- Format C (Spring 2025-2026 onward): one long "Data Entry" sheet for the
  whole quarter, with a "Week #" column.

load_qpass() detects the format, reads only the raw-data sheets, and returns
the same columns for every format. Values are kept raw (no renaming of
locations or types); cleaning happens in a later step.
"""

from __future__ import annotations

import datetime as dt
import logging
import numbers
import re
from collections.abc import Iterable
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from src.qpass_inventory import (
    HEADER_SEARCH_ROWS,
    QPASS_FILE_GLOB,
    RAW_QPASS_DIR,
    clean_column_name,
    parse_qpass_filename,
)

logger = logging.getLogger(__name__)

FORMAT_A = "A_wide_weekly"
FORMAT_B = "B_long_weekly"
FORMAT_C = "C_long_single"

# Format C keeps all data in one sheet with exactly this name.
FORMAT_C_SHEET = "Data Entry"
# Format B data sheets start with this prefix, e.g. "QPASS Data Entry (Week 1)".
FORMAT_B_SHEET_PREFIX = "QPASS Data Entry"

# Format A workbooks also contain formatted copies of the data (summaries,
# oversubscription views, re-sorted copies). Sheets whose names contain any of
# these words are not raw data and are skipped.
FORMAT_A_EXCLUDED_SHEET_WORDS = ("oversubscription", "summary", "sorted", "copy of", "current week")

# In Format A, the first column of the label row reads "Structure".
FORMAT_A_LABEL = "Structure"
FORMAT_A_LEVEL_PREFIX = "Level"
# Surface lots are named by ID: "P" followed by digits ("P002", "P785").
FORMAT_A_LOT_ID_PATTERN = r"P\d+"
# Text in the header row below the labels ("Available", "Available Spaces").
FORMAT_A_HEADER_ROW_TEXT = "Available"

# Bottom-of-sheet aggregate rows in some Format A sheets (e.g. Summer 2018).
# They hold totals across lots, not observations, so they are dropped.
# Compared without regard to case.
FORMAT_A_SUMMARY_ROW_LABELS = {
    "vacancy", "physical capacity", "occupancy", "oversubscription",
}

# Why a non-blank row inside a data sheet was not kept. Stored per sheet in
# attrs["dropped_rows"] as {reason: count}.
DROP_SUMMARY = "summary section (totals, not a parking location)"
DROP_NO_VALUES = "no counts and no total stall count"
DROP_NO_NAME = "no structure, lot or type"
DROP_NO_DATE = "no date"
DROP_INVALID_DATE = "invalid date"
# Formats B/C: rows the workbook template fills with formulas that point at an
# empty source row. Excel shows them as 0 / blank: location and type are 0 or
# blank and every count is 0 or blank.
DROP_PLACEHOLDER = "empty template row (no location, type or counts)"

# Excel stores a date as a number of days counted from this day (the 1900
# date system; 1899-12-30 rather than 12-31 because Excel treats 1900 as a
# leap year). Day 1 is the first valid Excel date.
EXCEL_EPOCH = pd.Timestamp("1899-12-30")
EXCEL_FIRST_SERIAL = 1

# QPASS counts open spaces at four times of day.
SURVEY_TIMES = {
    dt.time(8, 0): "open_8am",
    dt.time(10, 0): "open_10am",
    dt.time(12, 0): "open_12pm",
    dt.time(14, 0): "open_2pm",
}
OPEN_COLUMNS = list(SURVEY_TIMES.values())

# Raw header text (after clean_column_name) -> standardized column name,
# for Formats B and C. Only the first column with each header is used; later
# duplicates (e.g. a second "Day") are Excel-calculated helper columns.
LONG_FORMAT_COLUMNS = {
    "Date": "date",
    "Day": "day",
    "Week #": "week",
    "Zone": "zone",
    "Order": "source_order",
    "Report Area": "report_area",
    "Neighborhood": "neighborhood",
    "Lot / Structure": "lot_or_structure",
    "Location": "location",
    "Type": "type",
    "Sub Type": "sub_type",
    "Report Type": "report_type",
    "Total Stall Count": "total_stall_count",
    "Open Spaces 8am": "open_8am",
    "Open Spaces 10am": "open_10am",
    "Open Spaces 12pm": "open_12pm",
    "Open Spaces 2pm": "open_2pm",
}

# The header row of Format B/C data sheets starts with this cell.
LONG_FORMAT_HEADER_START = "Date"

# Columns a Format B/C data sheet must have. If a future workbook renames one,
# loading fails loudly instead of returning blank values.
LONG_FORMAT_REQUIRED_HEADERS = (
    "Date", "Location", "Type", "Total Stall Count",
    "Open Spaces 8am", "Open Spaces 10am", "Open Spaces 12pm", "Open Spaces 2pm",
)

# These columns must appear exactly once. Two "Open Spaces 8am" columns would
# mean two values for the same date and time, and one would be lost.
# Other repeated headers ("Day", "Sub Type") are Excel helper copies further
# right; the first one is used.
LONG_FORMAT_UNIQUE_HEADERS = LONG_FORMAT_REQUIRED_HEADERS

# Raw Format A label -> standardized column name.
FORMAT_A_COLUMNS = {
    "Structure": "raw_structure",
    "Lot": "raw_lot",
    "Type": "type",
    "Total Spaces": "total_stall_count",
}

# Every DataFrame returned by load_qpass has exactly these columns, in order.
OUTPUT_COLUMNS = [
    "source_file", "sheet_name", "source_row", "format_version",
    "academic_year", "quarter", "quarter_number",
    "date", "day", "week",
    "location", "lot_or_structure", "level",
    "zone", "source_order", "report_area", "neighborhood",
    "type", "sub_type", "report_type",
    "sub_type_in_format", "report_type_in_format",
    "total_stall_count", *OPEN_COLUMNS,
    "raw_structure", "raw_lot",
]

# Matches "Week 1" / "Week #2" / a standalone "W2". "W19" (= Winter 2019 in
# old sheet names) is not a week, so the short form allows one digit only.
WEEK_NUMBER_PATTERN = re.compile(r"Week\s*#?\s*(\d+)|\bW(\d)\b", re.IGNORECASE)


class UnsupportedQPASSFormatError(ValueError):
    """Raised when a workbook does not match any known QPASS layout."""


def load_qpass(path: str | Path) -> pd.DataFrame:
    """Load one QPASS workbook into a standardized long DataFrame.

    Each row is one location x space type x survey date, with open-space
    counts at 8am, 10am, 12pm and 2pm. academic_year, quarter and
    source_file come from the filename and are always filled. sheet_name and
    source_row (the Excel row number) point at the cells each row came from.

    Nothing is dropped silently. These are logged and stored on the
    returned DataFrame:

    - ``attrs["skipped_sheets"]``: list of (sheet name, reason) for sheets
      that were not read (non-data or empty).
    - ``attrs["dropped_rows"]``: {sheet name: {reason: count}} for non-blank
      rows inside data sheets that were not kept (see the DROP_* reasons).
    - ``attrs["non_numeric_values"]``: {column: count} of count or total
      cells holding text (e.g. "FULL", "-"). These rows are kept and the text
      is left as it is, for the cleaning step to decide.

    Args:
        path: Path to a QPASS .xlsx or .xlsm workbook.

    Returns:
        DataFrame with columns OUTPUT_COLUMNS. Optional *_in_format booleans
        record actual per-sheet header presence; source_order retains Order.

    Raises:
        FileNotFoundError: If the file does not exist.
        UnsupportedQPASSFormatError: If the filename does not follow the
            "QPASS {year}-{N} {Quarter}" pattern, no known data layout is
            found in the workbook, a data sheet is missing a required
            column, or two columns hold the same survey date and time.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    file_info = parse_qpass_filename(path.name)
    if file_info["academic_year"] is None:
        raise UnsupportedQPASSFormatError(
            f"{path.name}: filename does not match 'QPASS YYYY-YYYY-N Quarter'"
        )
    if not file_info["name_matches_number"]:
        raise UnsupportedQPASSFormatError(
            f"{path.name}: quarter number {file_info['quarter_number']} does not match "
            f"quarter name {file_info['quarter_name']!r}"
        )

    # By default pandas turns text such as "n/a", "NA" or "null" into blanks
    # while reading. Only truly empty cells should be blank, so that text
    # counts are kept as written.
    sheets = pd.read_excel(
        path, sheet_name=None, header=None, keep_default_na=False, na_values=[""]
    )
    format_version = detect_format(sheets)
    if format_version is None:
        raise UnsupportedQPASSFormatError(
            f"{path.name}: no '{FORMAT_C_SHEET}' sheet, no '{FORMAT_B_SHEET_PREFIX}...' "
            f"sheets, and no sheet with a '{FORMAT_A_LABEL}' label row"
        )

    frames, skipped, dropped = [], [], {}
    for sheet_name, sheet in sheets.items():
        try:
            if format_version == FORMAT_A:
                parsed, reason, sheet_dropped = _parse_format_a_sheet(sheet_name, sheet)
            else:
                parsed, reason, sheet_dropped = _parse_long_sheet(
                    sheet_name, sheet, format_version
                )
        except UnsupportedQPASSFormatError as error:
            raise UnsupportedQPASSFormatError(f"{path.name}: {error}") from error
        if parsed is None:
            skipped.append((sheet_name, reason))
            logger.info("%s: skipped sheet %r (%s)", path.name, sheet_name, reason)
            continue
        if sheet_dropped:
            dropped[sheet_name] = sheet_dropped
            for drop_reason, count in sheet_dropped.items():
                # An invalid date is unexpected, so it is a warning.
                is_invalid = drop_reason == DROP_INVALID_DATE
                level = logging.WARNING if is_invalid else logging.INFO
                logger.log(level, "%s: dropped %d rows in sheet %r: %s",
                           path.name, count, sheet_name, drop_reason)
        parsed["sheet_name"] = sheet_name
        # Record header presence before reindex fills absent columns with NaN.
        for column in ("sub_type", "report_type"):
            parsed[f"{column}_in_format"] = column in parsed.columns
        frames.append(parsed)

    if not frames:
        raise UnsupportedQPASSFormatError(
            f"{path.name}: detected {format_version} but no sheet contained data rows"
        )

    result = pd.concat(frames, ignore_index=True)
    result["source_file"] = path.name
    result["format_version"] = format_version
    result["academic_year"] = file_info["academic_year"]
    result["quarter"] = file_info["quarter_name"]
    result["quarter_number"] = file_info["quarter_number"]
    result = result.reindex(columns=OUTPUT_COLUMNS)
    _warn_on_dates_outside_academic_year(result, path.name)
    result.attrs["skipped_sheets"] = skipped
    result.attrs["dropped_rows"] = dropped
    result.attrs["non_numeric_values"] = _count_non_numeric_values(result, path.name)
    return result


def _count_non_numeric_values(data: pd.DataFrame, filename: str) -> dict[str, int]:
    """Count count/total cells that hold text instead of a number, and log them.

    Returns:
        {column name: count}, only for columns with at least one text value.
    """
    counts = {}
    for column in ["total_stall_count", *OPEN_COLUMNS]:
        values = data[column]
        is_text = values.notna() & pd.to_numeric(values, errors="coerce").isna()
        if is_text.any():
            counts[column] = int(is_text.sum())
            logger.info("%s: %d text values kept in %s (e.g. %r)",
                        filename, counts[column], column, values[is_text].iloc[0])
    return counts


def _collapse_spaces(name: str) -> str:
    """Return a file name with runs of whitespace collapsed to one space."""
    return " ".join(name.split())


def load_all_qpass(folder: str | Path = RAW_QPASS_DIR,
                   exclude: Iterable[str] = ()) -> pd.DataFrame:
    """Load every QPASS workbook in a folder into one combined DataFrame.

    This is the starting point for later steps (cleaning, name
    standardization, capacity matching): one table with the same columns
    for every year and format.

    pandas drops ``attrs`` when combining DataFrames, so each file's records
    are kept here, keyed by file name:

    - ``attrs["files"]``: names of the workbooks loaded, in order.
    - ``attrs["excluded_files"]``: workbooks present but not loaded because
      they are listed in ``exclude``.
    - ``attrs["skipped_sheets"]``: {file: [(sheet, reason), ...]}.
    - ``attrs["dropped_rows"]``: {file: {sheet: {reason: count}}}.
    - ``attrs["non_numeric_values"]``: {file: {column: count}}.

    Args:
        folder: Folder holding the workbooks. Defaults to data/raw/QPASS.
        exclude: Workbook file names to leave out (for example a workbook
            excluded by team decision). Names match with repeated spaces
            collapsed, so "QPASS 2024-2025-1  Summer.xlsm" matches either form.

    Returns:
        DataFrame with columns OUTPUT_COLUMNS; source_file, sheet_name and
        source_row say where each row came from.

    Raises:
        FileNotFoundError: If the folder has no QPASS workbooks to load.
        UnsupportedQPASSFormatError: If any workbook cannot be loaded. The
            error names the file; no file is skipped silently.
    """
    folder = Path(folder)
    excluded_names = {_collapse_spaces(name) for name in exclude}
    paths = sorted(folder.glob(QPASS_FILE_GLOB))
    excluded = [p.name for p in paths if _collapse_spaces(p.name) in excluded_names]
    paths = [p for p in paths if p.name not in excluded]
    for name in excluded:
        logger.info("%s: not loaded (excluded)", name)
    if not paths:
        raise FileNotFoundError(
            f"No files matching {QPASS_FILE_GLOB!r} in {folder}. "
            "Download them with: python -m src.download_qpass"
        )

    frames = {path.name: load_qpass(path) for path in paths}
    combined = pd.concat(frames.values(), ignore_index=True)
    combined.attrs = {"files": list(frames), "excluded_files": excluded}
    for key in ("skipped_sheets", "dropped_rows", "non_numeric_values"):
        combined.attrs[key] = {name: frame.attrs[key] for name, frame in frames.items()}
    logger.info("loaded %d QPASS workbooks, %d rows", len(paths), len(combined))
    return combined


def detect_format(sheets: dict[str, pd.DataFrame]) -> str | None:
    """Decide which QPASS layout a workbook uses, based on its sheets.

    Args:
        sheets: Mapping of sheet name to sheet contents (read with header=None).

    Returns:
        FORMAT_A, FORMAT_B, FORMAT_C, or None if no layout matches.
    """
    names = [name.strip() for name in sheets]
    if FORMAT_C_SHEET in names:
        return FORMAT_C
    if any(name.startswith(FORMAT_B_SHEET_PREFIX) for name in names):
        return FORMAT_B
    if any(_is_format_a_data_sheet(name, sheet) for name, sheet in sheets.items()):
        return FORMAT_A
    return None


def _find_row(sheet: pd.DataFrame, first_cell: str) -> int | None:
    """Return the first row (within HEADER_SEARCH_ROWS) whose first cell equals first_cell."""
    for row_index in range(min(HEADER_SEARCH_ROWS, len(sheet))):
        if clean_column_name(sheet.iat[row_index, 0]) == first_cell:
            return row_index
    return None


def _week_from_sheet_name(sheet_name: str) -> int | None:
    """Read a week number from names like "QPASS Data Entry (Week 1)" or "... W2"."""
    match = WEEK_NUMBER_PATTERN.search(sheet_name)
    if match is None:
        return None
    return int(match.group(1) or match.group(2))


def _parse_long_sheet(sheet_name: str, sheet: pd.DataFrame, format_version: str) -> tuple:
    """Parse a Format B or C data sheet.

    Returns:
        (DataFrame, None, dropped) on success, or (None, reason, {}) if the
        sheet is not a data sheet for this format or has no dated rows.
        dropped is {reason: count} for non-blank rows with a blank or
        invalid date, and for empty template rows (DROP_PLACEHOLDER).

    Raises:
        UnsupportedQPASSFormatError: If a required column is missing or a
            count column appears more than once.
    """
    name = sheet_name.strip()
    if format_version == FORMAT_C:
        is_data_sheet = name == FORMAT_C_SHEET
    else:
        is_data_sheet = name.startswith(FORMAT_B_SHEET_PREFIX)
    if not is_data_sheet:
        return None, "not a data sheet", {}

    header_row = _find_row(sheet, LONG_FORMAT_HEADER_START)
    if header_row is None:
        return None, f"no '{LONG_FORMAT_HEADER_START}' header row found", {}

    headers = [clean_column_name(value) for value in sheet.iloc[header_row]]
    for header in LONG_FORMAT_UNIQUE_HEADERS:
        columns = [
            get_column_letter(i + 1) for i, h in enumerate(headers) if h == header
        ]
        if len(columns) > 1:
            raise UnsupportedQPASSFormatError(
                f"sheet {sheet_name!r} has more than one {header!r} column "
                f"({', '.join(columns)}); cannot tell which value is real"
            )
    positions = {}
    for position, header in enumerate(headers):
        if header in LONG_FORMAT_COLUMNS and header not in positions:
            positions[header] = position
    _require_columns(sheet_name, positions, LONG_FORMAT_REQUIRED_HEADERS)

    body = sheet.iloc[header_row + 1:]
    data = pd.DataFrame({
        LONG_FORMAT_COLUMNS[header]: body.iloc[:, position].to_numpy()
        for header, position in positions.items()
    })
    is_non_blank = data.notna().any(axis=1)
    data["source_row"] = body.index + 1  # sheet row 0 is Excel row 1
    raw_date = data["date"]
    data["date"] = pd.to_datetime(raw_date.map(_parse_date_cell))
    has_date = data["date"].notna()
    counts_empty = data[OPEN_COLUMNS].apply(lambda col: col.map(_is_blank_or_zero))
    is_placeholder = (
        data["location"].map(_is_blank_or_zero)
        & data["type"].map(_is_blank_or_zero)
        & counts_empty.all(axis=1)
    )
    dropped = _nonzero({
        DROP_NO_DATE: (is_non_blank & raw_date.isna()).sum(),
        DROP_INVALID_DATE: (raw_date.notna() & ~has_date).sum(),
        DROP_PLACEHOLDER: (has_date & is_placeholder).sum(),
    })
    data = data[has_date & ~is_placeholder].reset_index(drop=True)
    if data.empty:
        return None, "no rows with a date (empty template sheet)", {}

    if "week" in data:
        data["week"] = data["week"].map(_parse_week_value)
    else:
        data["week"] = _week_from_sheet_name(name)
    return data, None, dropped


def _is_blank_or_zero(value) -> bool:
    """True for a blank cell or a zero, the values Excel shows for an empty formula."""
    if not _is_filled(value):
        return True
    return pd.to_numeric(value, errors="coerce") == 0


def _parse_date_cell(value) -> pd.Timestamp:
    """Turn one Date cell into a Timestamp, or NaT if it is not a valid date.

    A number is an Excel serial date (days since EXCEL_EPOCH). Without this
    step pandas would read 45565 as nanoseconds after 1970-01-01, giving a
    false 1970 date instead of 2024-09-30.
    """
    if pd.isna(value):
        return pd.NaT
    if isinstance(value, numbers.Real) and not isinstance(value, bool):
        if value < EXCEL_FIRST_SERIAL:
            return pd.NaT
        try:
            return (EXCEL_EPOCH + pd.Timedelta(days=float(value))).normalize()
        except (OverflowError, ValueError):  # too large to be a date
            return pd.NaT
    return pd.to_datetime(value, errors="coerce")


def _nonzero(counts: dict) -> dict[str, int]:
    """Keep only the reasons with a count above zero, as plain ints."""
    return {reason: int(count) for reason, count in counts.items() if count}


def _require_columns(sheet_name: str, found: dict, required: tuple) -> None:
    """Raise UnsupportedQPASSFormatError if any required header is missing.

    Args:
        sheet_name: Sheet being parsed, used in the error message.
        found: Headers found in the sheet (only the keys are used).
        required: Headers the sheet must contain.
    """
    missing = [header for header in required if header not in found]
    if missing:
        raise UnsupportedQPASSFormatError(
            f"sheet {sheet_name!r} is missing required column(s): {', '.join(missing)}"
        )


def _parse_week_value(value) -> int | None:
    """Turn a "Week #" cell ("Week 0", 1, "1") into an int, or None if blank."""
    if pd.isna(value):
        return None
    digits = re.search(r"\d+", str(value))
    return int(digits.group()) if digits else None


def _is_format_a_data_sheet(sheet_name: str, sheet: pd.DataFrame) -> bool:
    """True if a sheet is a raw Format A weekly data sheet."""
    lowered = sheet_name.lower()
    if any(word in lowered for word in FORMAT_A_EXCLUDED_SHEET_WORDS):
        return False
    label_row = _find_row(sheet, FORMAT_A_LABEL)
    return label_row is not None and bool(_format_a_column_dates(sheet, label_row))


def _format_a_column_dates(sheet: pd.DataFrame, label_row: int) -> dict[int, tuple]:
    """Map each count column of a Format A sheet to its (date, survey time).

    The header block above the data has a row of dates (one per day block)
    and a row of times (one per column). A column's date is the nearest date
    at or to its left, because each date sits above the first of its four
    time columns.

    Returns:
        {column_index: (date, open_column_name)} for every survey-time column.
    """
    header_block = sheet.iloc[: label_row + 1]
    column_dates = {}
    current_date = None
    for column in range(sheet.shape[1]):
        cells = header_block.iloc[:, column].tolist()
        for cell in cells:
            if isinstance(cell, (dt.datetime, pd.Timestamp)):
                current_date = pd.Timestamp(cell).normalize()
        survey_time = next((cell for cell in cells if isinstance(cell, dt.time)), None)
        if current_date is not None and survey_time in SURVEY_TIMES:
            column_dates[column] = (current_date, SURVEY_TIMES[survey_time])
    return column_dates


def _parse_format_a_sheet(sheet_name: str, sheet: pd.DataFrame) -> tuple:
    """Parse one Format A weekly sheet into long rows (one per row x date).

    Structure and Lot cells are often blank, meaning "same as the row above".
    Lot is filled down, but never past a row that names a new Structure.
    Structure is filled down only for level rows ("Level 2") and for structure
    rows with no lot at all; surface lots ("P002") with a blank Structure do
    not belong to the structure above them.

    Summary sections are dropped. A summary section starts at a row labelled
    with a FORMAT_A_SUMMARY_ROW_LABELS word in Structure or Lot (Summer 2018:
    "A" / "Vacancy") or at a heading-only Structure row such as
    "West Campus" (2019 files, followed by subtotal rows like "A" 1593)
    whose rows are not levels or lot IDs. It lasts until the next row with
    a non-blank Structure. See _format_a_summary_rows.

    Rows are kept if they name a structure, lot or type and have any count or
    total, even text such as "FULL" or "-" (kept as written, like Formats
    B and C).

    Returns:
        (DataFrame, None, dropped) on success, or (None, reason, {}) if not
        a data sheet. dropped is {reason: count} for non-blank rows that were
        not kept.

    Raises:
        UnsupportedQPASSFormatError: If a required label is missing, or two
            columns hold the same survey date and time.
    """
    if not _is_format_a_data_sheet(sheet_name, sheet):
        return None, "not a data sheet", {}

    label_row = _find_row(sheet, FORMAT_A_LABEL)
    labels = [clean_column_name(value) for value in sheet.iloc[label_row]]
    label_positions = {label: labels.index(label) for label in FORMAT_A_COLUMNS if label in labels}
    _require_columns(sheet_name, label_positions, tuple(FORMAT_A_COLUMNS))
    column_dates = _format_a_column_dates(sheet, label_row)
    _reject_duplicate_date_times(sheet_name, column_dates)

    body = sheet.iloc[label_row + 1:]
    # Some sheets have an "Available Available ..." row under the labels. It
    # is part of the header, so remove it before counting dropped rows.
    is_header_row = body.apply(
        lambda row: row.dropna().astype(str).str.startswith(FORMAT_A_HEADER_ROW_TEXT).all()
        and row.notna().any(),
        axis=1,
    )
    body = body[~is_header_row]
    source_row = body.index + 1  # sheet row 0 is Excel row 1
    body = body.reset_index(drop=True)
    info = pd.DataFrame({
        FORMAT_A_COLUMNS[label]: body.iloc[:, position].to_numpy()
        for label, position in label_positions.items()
    })

    is_filled = body.apply(lambda column: column.map(_is_filled))
    has_values = is_filled.iloc[:, list(column_dates)].any(axis=1)
    has_values |= info["total_stall_count"].map(_is_filled)
    structure_text = info["raw_structure"].map(clean_column_name)
    lot_text = info["raw_lot"].map(clean_column_name)
    has_name = (structure_text != "") | (lot_text != "") | info["type"].map(_is_filled)
    structure_position = label_positions[FORMAT_A_LABEL]
    other_cells = is_filled.drop(columns=body.columns[structure_position])
    is_structure_only = ~other_cells.any(axis=1)
    has_type = info["type"].map(_is_filled)
    is_summary = _format_a_summary_rows(
        structure_text, lot_text, is_structure_only, has_values, has_type
    )

    structure = structure_text.replace("", None)
    filled_structure = structure.ffill()
    # Each non-blank Structure cell starts a new group. Lot is filled down only
    # within its group, so a structure listed with a blank Lot (e.g. "Gilman")
    # does not inherit the surface lot above it (e.g. "P010").
    structure_group = structure.notna().cumsum()
    lot = lot_text.replace("", None).groupby(structure_group).ffill()
    is_level = lot.fillna("").str.startswith(FORMAT_A_LEVEL_PREFIX)
    is_whole_structure = lot.isna()  # structure row with no lot or level given
    info["location"] = filled_structure.where(is_level | is_whole_structure, lot)
    info["level"] = lot.where(is_level)
    info["lot_or_structure"] = (is_level | is_whole_structure).map({True: "Structure", False: "Lot"})
    info["source_row"] = source_row.to_numpy()

    keep = has_values & has_name & ~is_summary
    not_kept = is_filled.any(axis=1) & ~keep
    dropped = _nonzero({
        DROP_NO_VALUES: (not_kept & ~has_values).sum(),
        DROP_SUMMARY: (not_kept & has_values & is_summary).sum(),
        DROP_NO_NAME: (not_kept & has_values & ~is_summary & ~has_name).sum(),
    })
    info, body = info[keep], body[keep]
    if info.empty:
        return None, "no data rows below the header", {}

    records = []
    dates = sorted({date for date, _ in column_dates.values()})
    for date in dates:
        day_frame = info.copy()
        day_frame["date"] = date
        day_frame["day"] = date.day_name()
        for column, (column_date, open_name) in column_dates.items():
            if column_date == date:
                day_frame[open_name] = body.iloc[:, column].to_numpy()
        records.append(day_frame)
    data = pd.concat(records, ignore_index=True)
    data["week"] = _week_from_sheet_name(sheet_name)
    return data, None, dropped


def _is_filled(value) -> bool:
    """True if a cell holds something other than blank or whitespace."""
    if isinstance(value, str):
        return bool(value.strip())
    return not pd.isna(value)


def _format_a_summary_rows(
    structure_text: pd.Series,
    lot_text: pd.Series,
    is_structure_only: pd.Series,
    has_values: pd.Series,
    has_type: pd.Series,
) -> pd.Series:
    """Mark the rows of a Format A sheet that belong to a summary section.

    A section starts at every row with a non-blank Structure and at every
    row whose Lot is a summary label. The whole section is a summary if its
    first row:

    - has a summary label ("Vacancy"), or
    - is a heading, a row where only the Structure cell is filled
      ("West Campus", "SIO"), and none of the section's rows with values
      is a real observation. A real observation has a level in Lot
      ("Level 1"), a lot ID ("P002"), or a blank Lot with a Type (a
      structure without levels). Under the real headings every row's Lot
      is a space type ("A", "SIO B"), but a structure name on its own row
      followed by its levels ("Athena", then "Level 1") is not a summary.

    Returns:
        Boolean Series, True for rows inside a summary section.
    """
    is_label = (
        structure_text.str.casefold().isin(FORMAT_A_SUMMARY_ROW_LABELS)
        | lot_text.str.casefold().isin(FORMAT_A_SUMMARY_ROW_LABELS)
    )
    starts_section = (structure_text != "") | is_label
    section = starts_section.cumsum()
    is_observation = has_values & (
        lot_text.str.startswith(FORMAT_A_LEVEL_PREFIX)
        | lot_text.str.fullmatch(FORMAT_A_LOT_ID_PATTERN)
        | ((lot_text == "") & has_type)
    )
    has_observations = is_observation.groupby(section).transform("any")
    is_heading = (structure_text != "") & is_structure_only & ~has_observations
    starts_summary = starts_section & (is_label | is_heading)
    return starts_summary.groupby(section).transform("any")


def _reject_duplicate_date_times(
    sheet_name: str, column_dates: dict[int, tuple]
) -> None:
    """Raise if two Format A columns map to the same survey date and time.

    Otherwise one column's values would silently overwrite the other's.
    """
    columns_by_slot = {}
    for column, slot in column_dates.items():
        columns_by_slot.setdefault(slot, []).append(get_column_letter(column + 1))
    for (date, open_name), columns in columns_by_slot.items():
        if len(columns) > 1:
            raise UnsupportedQPASSFormatError(
                f"sheet {sheet_name!r}: columns {', '.join(columns)} are all "
                f"{date.date()} {open_name}; cannot tell which value is real"
            )


def _warn_on_dates_outside_academic_year(data: pd.DataFrame, filename: str) -> None:
    """Log a warning if any row's date falls outside the file's academic year.

    An academic year "2024-2025" spans calendar years 2024 and 2025.
    """
    start_year, end_year = (int(year) for year in data["academic_year"].iloc[0].split("-"))
    outside = ~data["date"].dt.year.between(start_year, end_year)
    if outside.any():
        logger.warning(
            "%s: %d rows have dates outside academic year %d-%d",
            filename, int(outside.sum()), start_year, end_year,
        )
