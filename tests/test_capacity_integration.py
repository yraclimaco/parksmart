"""Run real Excel parsing, naming, extraction and capacity selection together."""

import datetime as dt
import json

import pandas as pd
import pytest
from openpyxl import Workbook

from src.compare_capacities import compare_capacities
from src.clean_inventory_modern import clean_modern_inventory
from src.run_capacity_matching import run_capacity_matching


def write_qpass(path, rows, date=dt.datetime(2024, 9, 30)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QPASS Data Entry (Week 1)"
    sheet.append(["Name(s):"])
    sheet.append(["Date", "Day", "Zone", "Order", "Report Area", "Neighborhood",
                  "Lot / Structure", "Location", "Type", "Report Type",
                  "Total Stall Count", "Open Spaces 8am", "Open Spaces 10am",
                  "Open Spaces 12pm", "Open Spaces 2pm"])
    for facility, category, capacity, readings in rows:
        sheet.append([date, "Monday", 1, 1,
                      "Main Campus", "Central", "Lot", facility, category,
                      category, capacity, *readings])
    workbook.save(path)
    workbook.close()


def write_inventory(path):
    workbook = Workbook()
    workbook.active.title = "Title"
    workbook.active.append(["Parking Space Inventory Fall 2024"])
    sheet = workbook.create_sheet("Pivot Table Data")
    sheet.append(["Lot / Garage", "Space Type", "Space Subtype",
                  "Time Limit [minutes]", "Inventory", "Fall 2024"])
    sheet.append(["P386", "Visitor", None, None, 10, 10])
    workbook.save(path)
    workbook.close()


def test_real_excel_loader_labels_cleaner_and_batch(tmp_path):
    qpass_dir = tmp_path / "QPASS"
    qpass_dir.mkdir()
    qpass_path = qpass_dir / "QPASS 2024-2025-2 Fall.xlsx"
    write_qpass(qpass_path, [
        (" Gliderport ", "V", None, [None, 6, 4, 2]),
        ("P100", "B", 20, [10, 9, 8, 7]),
        (None, "B", 12, [1, 1, 1, 1]),
    ])
    inventory = tmp_path / "Inventory 2024-2025-2 Fall.xlsx"
    write_inventory(inventory)
    manifest = tmp_path / "manifest.csv"
    pd.DataFrame([dict(academic_year="2024-2025", quarter="Fall",
                       qpass_file=qpass_path.name, inventory_path=str(inventory),
                       download_status="already_present", reason="")]).to_csv(manifest, index=False)
    quarter, total = run_capacity_matching(
        manifest, qpass_dir, tmp_path / "clean", tmp_path / "reports",
        rebuild_inventory=True,
    )
    assert total["total_observation_slots"] == 12
    assert total["eligible_observation_slots"] == 7
    assert total["excluded_observation_slots"] == 5
    assert total["accepted_qpass_capacities"] == 4
    assert total["accepted_inventory_capacities"] == 3
    assert total["accepted_capacity_match_rate"] == 1
    for field in total:
        if field != "qpass_quarters":
            assert quarter.iloc[0][field] == pytest.approx(total[field])
    result = pd.read_csv(tmp_path / "reports/combined/capacity_comparison.csv",
                         keep_default_na=False)
    fallback = result.loc[result.capacity_source.eq("inventory_same_quarter")]
    assert fallback.facility.eq("P386").all()
    assert fallback.category.eq("Visitor").all()
    assert fallback.qpass_facility_raw.eq(" Gliderport ").all()
    assert fallback.qpass_category_raw.eq("V").all()
    assert fallback.capacity_assumption.ne("").all()
    assert pd.to_numeric(fallback.selected_capacity).eq(10).all()
    assert fallback.qpass_capacity_raw.eq("").all()
    assert result.loc[result.qpass_source_row.eq(5), "exclusion_reason"].eq(
        "invalid_facility_or_category"
    ).all()
    cleaned = pd.read_csv(tmp_path / "reports/inventory_cleaned.csv")
    raw_group = json.loads(cleaned.iloc[0].source_groups_json)[0]
    assert raw_group["category_raw"] == "Visitor"
    assert inventory.is_file()


@pytest.mark.parametrize("field", ["location", "type"])
@pytest.mark.parametrize("value", [None, float("nan"), "nan", "", "0", "0.0"])
def test_missing_ids_excluded_from_denominators(tmp_path, monkeypatch, field, value):
    import src.compare_capacities as comparison
    source = dict(academic_year="2024-2025", quarter="Fall", location="P386", type="V",
                  sub_type=None, level=None, date=dt.datetime(2024, 9, 30),
                  total_stall_count=10, source_file="test.xlsx", sheet_name="Week 1",
                  source_row=3, open_8am=1, open_10am=2, open_12pm=3, open_2pm=4)
    source[field] = value
    monkeypatch.setattr(comparison, "load_qpass", lambda _: pd.DataFrame([source]))
    inv = tmp_path / "Inventory 2024-2025-2 Fall.xlsx"
    write_inventory(inv)
    csv = clean_modern_inventory(inv, "2024-2025", "Fall", tmp_path / "clean")
    result = compare_capacities(csv, tmp_path / "test.xlsx", tmp_path / "out")
    assert len(result) == 4
    assert result.exclusion_reason.eq("invalid_facility_or_category").all()
    total = pd.read_csv(tmp_path / "out/overall_summary.csv").iloc[0]
    assert total.total_observation_slots == 4
    assert total.eligible_observation_slots == 0
    assert pd.isna(total.accepted_capacity_match_rate)


def test_known_formula_problem_remains_source_data_issue(tmp_path, monkeypatch):
    import src.compare_capacities as comparison
    source = dict(academic_year="2025-2026", quarter="Fall", location="P100", type="B",
                  sub_type=None, level=None, date=dt.datetime(2025, 9, 29),
                  total_stall_count=0, source_file="QPASS 2025-2026-2 Fall.xlsm",
                  sheet_name="Week 1", source_row=3,
                  open_8am=1, open_10am=2, open_12pm=3, open_2pm=4)
    monkeypatch.setattr(comparison, "load_qpass", lambda _: pd.DataFrame([source]))
    inventory = tmp_path / "inventory.csv"
    pd.DataFrame([dict(academic_year="2025-2026", quarter="Fall", facility_raw="P100",
                       category_raw="B", capacity="", capacity_status="period_metadata_review",
                       source_file="Inventory.xlsm", source_rows="5")]).to_csv(inventory, index=False)
    result = compare_capacities(inventory, tmp_path / "qpass.xlsm", tmp_path / "out")
    assert result.qpass_capacity_raw.eq(0).all()
    assert result.selected_capacity.isna().all()
    assert result.unresolved_type.eq("source_data_problem").all()
    assert result.qpass_source_issue.eq("fall_2025_zero_capacity_formula").all()


def test_two_quarter_rebuild_and_deliverable_validation(tmp_path, monkeypatch):
    """Legacy and modern workbooks rebuild, join and package with raw evidence."""
    import importlib.util
    from pathlib import Path
    from src.clean_inventory_legacy import CATEGORY_HEADERS

    qpass_dir = tmp_path / "data/raw/QPASS"
    qpass_dir.mkdir(parents=True)
    raw_dir = tmp_path / "data/raw/inventory"
    raw_dir.mkdir()
    (tmp_path / "docs").mkdir()
    source_dir = tmp_path / "results/capacity_matching"
    source_dir.mkdir(parents=True)
    modern = raw_dir / "Inventory 2024-2025-2 Fall.xlsx"
    write_inventory(modern)
    modern_qpass = qpass_dir / "QPASS 2024-2025-2 Fall.xlsx"
    write_qpass(modern_qpass, [("Gliderport", "V", None, [None, 6, 4, 2])])
    legacy = raw_dir / "Inventory 2018-2019-2 October 1.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, first, facility in [("5A By Lot", 2, "P100"),
                                   ("6A By Structure", 1, "Hopkins")]:
        sheet = workbook.create_sheet(name)
        rows = [[None] * 15 for _ in range(8)]
        rows[1][first - 1] = "Inventory Fall 2018"
        rows[5][first:first + 10] = CATEGORY_HEADERS
        rows[6][first + 7] = "Vehicle"
        rows[6][first + 8] = "Yard"
        rows[7][first - 1] = facility
        rows[7][first] = 25
        rows[7][first + 7] = 4
        for row in rows:
            sheet.append(row)
    workbook.save(legacy)
    workbook.close()
    legacy_qpass = qpass_dir / "QPASS 2018-2019-2 Fall.xlsx"
    write_qpass(legacy_qpass, [("P100", "A", None, [10, 9, 8, 7]),
                              ("P100", "UC Vehicle", 4, [1, 1, 1, 1])],
                date=dt.datetime(2018, 10, 4))
    records = [dict(academic_year=year, quarter="Fall", qpass_file=qpass.name,
                    inventory_path=str(inv), download_status="already_present", reason="")
               for year, qpass, inv in [("2024-2025", modern_qpass, modern),
                                         ("2018-2019", legacy_qpass, legacy)]]
    pd.DataFrame(records).to_csv(source_dir / "inventory_sources.csv", index=False)
    before = {p: p.read_bytes() for p in (modern, legacy, modern_qpass, legacy_qpass)}
    run_capacity_matching(source_dir / "inventory_sources.csv", qpass_dir,
                          tmp_path / "data/interim/inventory", source_dir / "batch",
                          rebuild_inventory=True)
    script = Path(__file__).resolve().parents[1] / "scripts/prepare_capacity_deliverables.py"
    spec = importlib.util.spec_from_file_location("prepare_deliverables", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    validation = module.prepare(tmp_path)
    assert validation["checks_passed"]
    assert validation["qpass_quarters"] == 2
    assert validation["inventory_workbooks"] == 2
    assert validation["counts"]["accepted_inventory_capacities"] == 3
    assert validation["counts"]["accepted_qpass_capacities"] == 4
    for path, original in before.items():
        assert path.read_bytes() == original
    observations = pd.read_csv(source_dir / "batch/combined/capacity_comparison.csv")
    assert observations.loc[observations.qpass_category_raw.eq("UC Vehicle"),
                            "category"].eq("UC Vehicle").all()
    assert (source_dir / "deliverables/capacity_comparison.csv.gz").is_file()
    # A tampered report must fail validation before replacing deliveries.
    overall_path = source_dir / "batch/overall_summary.csv"
    overall = pd.read_csv(overall_path)
    overall.loc[0, "accepted_capacity_matches"] += 1
    overall.to_csv(overall_path, index=False)
    with pytest.raises(ValueError, match="Count mismatch"):
        module.prepare(tmp_path)


@pytest.mark.parametrize("subtype,limit,level", [
    ("Restricted", "", ""), ("", "120", ""), ("", "", "Level 1")
])
def test_broad_inventory_cannot_fill_restricted_qpass(tmp_path, monkeypatch,
                                                     subtype, limit, level):
    import src.compare_capacities as comparison
    source = dict(academic_year="2025-2026", quarter="Fall", location="P100", type="Visitor",
                  sub_type=subtype, level=level, date=dt.datetime(2025, 9, 29),
                  total_stall_count=None, source_file="QPASS 2025-2026-2 Fall.xlsm",
                  sheet_name="Week 1", source_row=3,
                  open_8am=1, open_10am=2, open_12pm=3, open_2pm=4)
    monkeypatch.setattr(comparison, "load_qpass", lambda _: pd.DataFrame([source]))
    inventory = tmp_path / "inventory.csv"
    pd.DataFrame([dict(academic_year="2025-2026", quarter="Fall", facility_raw="P100",
                       category_raw="Visitor", subtype_raw="", time_limit_raw=limit,
                       capacity=10, capacity_status="positive_unverified",
                       source_file="Inventory.xlsm", source_rows="5")]).to_csv(inventory, index=False)
    result = compare_capacities(inventory, tmp_path / "qpass.xlsm", tmp_path / "out")
    assert result.selected_capacity.isna().all()
    assert result.unresolved_type.eq("scope_conflict").all()


def test_known_formula_zero_can_fallback_and_remains_flagged(tmp_path, monkeypatch):
    import src.compare_capacities as comparison
    source = dict(academic_year="2025-2026", quarter="Fall", location="P100", type="B",
                  sub_type=None, level=None, date=dt.datetime(2025, 9, 29),
                  total_stall_count=0, source_file="QPASS 2025-2026-2 Fall.xlsm",
                  sheet_name="Week 1", source_row=3,
                  open_8am=1, open_10am=2, open_12pm=3, open_2pm=4)
    monkeypatch.setattr(comparison, "load_qpass", lambda _: pd.DataFrame([source]))
    inventory = tmp_path / "inventory.csv"
    pd.DataFrame([dict(academic_year="2025-2026", quarter="Fall", facility_raw="P100",
                       category_raw="B", capacity=10, capacity_status="positive_unverified",
                       source_file="Inventory.xlsm", source_rows="5")]).to_csv(inventory, index=False)
    result = compare_capacities(inventory, tmp_path / "qpass.xlsm", tmp_path / "out")
    assert result.selected_capacity.eq(10).all()
    assert result.capacity_source.eq("inventory_same_quarter").all()
    assert result.qpass_capacity_raw.eq(0).all()
    assert result.qpass_source_issue.ne("").all()


def test_mapping_restriction_and_period_issue_stay_distinct(tmp_path, monkeypatch):
    import src.compare_capacities as comparison
    base = dict(academic_year="2025-2026", quarter="Fall", location="P100", type="B",
                sub_type=None, level=None, date=dt.datetime(2025, 9, 29),
                total_stall_count=None, source_file="QPASS 2025-2026-2 Fall.xlsm",
                sheet_name="Week 1", source_row=3,
                open_8am=1, open_10am=2, open_12pm=3, open_2pm=4)
    rows = [base, {**base, "type": "24/7", "source_row": 4},
            {**base, "type": "2 Person Carpool", "source_row": 5, "total_stall_count": 10}]
    monkeypatch.setattr(comparison, "load_qpass", lambda _: pd.DataFrame(rows))
    inventory = tmp_path / "inventory.csv"
    pd.DataFrame([dict(academic_year="2025-2026", quarter="Fall", facility_raw="P100",
                       category_raw="B", capacity="", capacity_status="period_metadata_review",
                       source_file="Inventory.xlsm", source_rows="5")]).to_csv(inventory, index=False)
    result = compare_capacities(inventory, tmp_path / "qpass.xlsm", tmp_path / "out")
    assert result.loc[result.qpass_source_row.eq(3), "unresolved_type"].eq("period_metadata_conflict").all()
    assert result.loc[result.qpass_source_row.eq(4), "unresolved_type"].eq("mapping_review").all()
    restricted = result.loc[result.qpass_source_row.eq(5)]
    assert restricted.category.eq("Carpool 2 Person").all()
    assert restricted.qpass_category_status.eq("keep_separate").all()
    assert restricted.qpass_category_raw.eq("2 Person Carpool").all()
    assert restricted.selected_capacity.eq(10).all()


@pytest.mark.parametrize("reported", [None, 10])
def test_wrong_year_real_excel_cannot_select_capacity(tmp_path, reported):
    qpass = tmp_path / "QPASS 2024-2025-2 Fall.xlsx"
    write_qpass(qpass, [("P386", "Visitor", reported, [1, 2, 3, 4])],
                date=dt.datetime(1900, 1, 1))
    inventory = tmp_path / "Inventory 2024-2025-2 Fall.xlsx"
    write_inventory(inventory)
    groups = clean_modern_inventory(inventory, "2024-2025", "Fall", tmp_path / "clean")
    result = compare_capacities(groups, qpass, tmp_path / "out")
    assert len(result) == 4
    assert result.selected_capacity.isna().all()
    assert result.unresolved_type.eq("period_metadata_conflict").all()
    assert result.qpass_period_metadata_conflict.all()
    assert pd.to_datetime(result.observation_date).eq(pd.Timestamp("1900-01-01")).all()
    assert len(pd.read_csv(tmp_path / "out/unresolved_cases.csv")) == 4


@pytest.mark.parametrize("reuse", [False, True])
def test_all_missing_inventory_outputs_are_empty_not_stale(tmp_path, reuse):
    qpass_dir = tmp_path / "QPASS"
    qpass_dir.mkdir()
    qpass = qpass_dir / "QPASS 2024-2025-2 Fall.xlsx"
    write_qpass(qpass, [("P386", "Visitor", 10, [1, 2, 3, 4])])
    manifest = tmp_path / "manifest.csv"
    pd.DataFrame([dict(academic_year="2024-2025", quarter="Fall",
                       qpass_file=qpass.name, inventory_path="",
                       download_status="not_listed_in_archive", reason="Missing")]
                 ).to_csv(manifest, index=False)
    output = tmp_path / "out"
    output.mkdir()
    if reuse:
        for name in ("inventory_cleaned.csv", "inventory_capacity_groups.csv"):
            (output / name).write_text("stale,value\nold,999\n")
    run_capacity_matching(manifest, qpass_dir, tmp_path / "clean", output)
    for name in ("inventory_cleaned.csv", "inventory_capacity_groups.csv"):
        groups = pd.read_csv(output / name)
        assert groups.empty
        assert {"academic_year", "quarter", "facility_raw", "capacity"} <= set(groups)
    assert len(pd.read_csv(output / "combined/matched_records.csv")) == 4


def test_deprecated_audit_requires_period_and_keeps_invalid_identifiers(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from src.audit_inventory import audit_inventory

    inventory = tmp_path / "Inventory 2024-2025-2 Fall.xlsx"
    write_inventory(inventory)
    workbook = load_workbook(inventory)
    workbook["Pivot Table Data"].append([None, "B", None, None, 10, 10])
    workbook.save(inventory)
    workbook.close()
    before = inventory.read_bytes()
    monkeypatch.chdir(tmp_path)
    with pytest.raises(ValueError, match="academic-year"):
        audit_inventory(inventory)
    with pytest.warns(FutureWarning, match="deprecated"):
        path = audit_inventory(inventory, "2024-2025", "Fall")
    groups = pd.read_csv(path)
    assert len(groups) == 2
    assert groups.capacity_status.eq("invalid_identifier_review").sum() == 1
    assert inventory.read_bytes() == before


@pytest.mark.parametrize("reported,verified,evidence,expected", [
    (30, False, "", 30),
    (None, False, "", None),
    (0, False, "", None),
    (None, True, "", None),
    (30, True, "Reviewed exact group in source cells", None),
    (None, True, "Reviewed exact group in source cells", 25),
])
def test_legacy_summary_requires_explicit_scope_verification(
    tmp_path, reported, verified, evidence, expected
):
    """Unverified summaries neither contradict nor fill; verified conflicts remain."""
    inventory = tmp_path / "legacy.csv"
    pd.DataFrame([dict(
        academic_year="2018-2019", quarter="Fall", facility_raw="P100",
        category_raw="A", category="A", capacity=25,
        capacity_status="positive_unverified", source_file="Inventory legacy.xlsx",
        source_rows="8", schema_version="legacy_summary_tables_v1",
        category_scope="legacy_summary_category_requires_review",
        scope_verified=verified, scope_verification_evidence=evidence,
    )]).to_csv(inventory, index=False)
    qpass = tmp_path / "QPASS 2018-2019-2 Fall.xlsx"
    write_qpass(qpass, [("P100", "A", reported, [5, 4, 3, 2])],
                date=dt.datetime(2018, 10, 4))
    result = compare_capacities(inventory, qpass, tmp_path / "out")
    assert len(result) == 4
    if expected is None:
        assert result.selected_capacity.isna().all()
    else:
        assert result.selected_capacity.eq(expected).all()
    assert result.candidate_capacity.eq(25).all()
    assert result.inventory_source_groups_json.str.contains(
        "legacy_summary_category_requires_review"
    ).all()
    if not verified or not evidence:
        assert result.inventory_summary_scope_unverified.all()
        assert not result.inventory_scope_compatible.any()
        assert not result.group_capacity_conflict.any()
        assert result.capacity_difference.isna().all()
        assert result.reported_capacity_comparison.eq("scope_review_required").all()
    elif reported == 30:
        assert result.group_capacity_conflict.all()
        assert result.reported_capacity_comparison.eq("conflict").all()


def test_legacy_schema_without_scope_marker_is_not_implicitly_verified(tmp_path):
    inventory = tmp_path / "legacy.csv"
    pd.DataFrame([dict(academic_year="2018-2019", quarter="Fall", facility_raw="P100",
                       category_raw="A", capacity=25, capacity_status="positive_unverified",
                       source_file="legacy.xlsx", source_rows="8",
                       schema_version="legacy_summary_tables_v1")]).to_csv(inventory, index=False)
    qpass = tmp_path / "QPASS 2018-2019-2 Fall.xlsx"
    write_qpass(qpass, [("P100", "A", 30, [5, 4, 3, 2])], date=dt.datetime(2018, 10, 4))
    result = compare_capacities(inventory, qpass, tmp_path / "out")
    assert result.selected_capacity.eq(30).all()
    assert result.inventory_summary_scope_unverified.all()


def test_documented_commands_use_fresh_downloader_filename(tmp_path, monkeypatch):
    """Execute both docs' commands against a newly downloaded canonical file."""
    import io
    import os
    from pathlib import Path
    import shlex
    import subprocess
    import sys
    from unittest.mock import patch
    from src import download_inventory as downloader

    project = Path(__file__).resolve().parents[1]
    qpass_dir = tmp_path / "data/raw/QPASS"
    qpass_dir.mkdir(parents=True)
    write_qpass(qpass_dir / "QPASS 2025-2026-2 Fall.xlsm",
                [("P386", "Visitor", None, [1, 2, 3, 4])],
                date=dt.datetime(2025, 9, 29))
    workbook = Workbook()
    workbook.active.title = "Title"
    workbook.active.append(["Parking Space Inventory Fall 2025"])
    sheet = workbook.create_sheet("Pivot Table Data")
    sheet.append(["Lot  / Structure", "Space Type", "Space Subtype",
                  "Time Limit [minutes]", "Inventory By Location"])
    sheet.append(["P386", "Visitor", None, None, 10])
    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()
    html = ('<table><tr><td>2025/2026</td><td></td><td><a href="'
            'Inventory%202025-2026-2%20Fall.xlsm">Excel</a></td>'
            '<td></td><td></td></tr></table>')
    monkeypatch.chdir(tmp_path)
    with patch.object(downloader, "urlopen", side_effect=[
        io.BytesIO(html.encode()), io.BytesIO(payload.getvalue()),
    ]):
        records = downloader.download_inventory()
    assert records[0]["download_status"] == "downloaded"
    source = tmp_path / "data/raw/inventory/Inventory 2025-2026-2 Fall.xlsm"
    before = source.read_bytes()
    env = {**os.environ, "PYTHONPATH": str(project)}
    for doc in ("AGENTS.md", "docs/inventory_audit.md"):
        commands = [line for line in (project / doc).read_text().splitlines()
                    if line.startswith("python -m src.clean_inventory_modern ")
                    or line.startswith("python -m src.compare_capacities ")]
        assert len(commands) == 2
        for command in commands:
            assert "Inventory_2025-2026_Fall" not in command
            subprocess.run([sys.executable, *shlex.split(command)[1:]],
                           cwd=tmp_path, env=env, capture_output=True, text=True,
                           check=True)
    groups = tmp_path / "data/interim/inventory/Inventory 2025-2026-2 Fall_capacity_groups.csv"
    assert groups.is_file()
    results = pd.read_csv(tmp_path / "results/capacity_matching/fall_2025/capacity_comparison.csv")
    assert len(results) == 4
    assert results.selected_capacity.eq(10).all()
    assert source.read_bytes() == before
