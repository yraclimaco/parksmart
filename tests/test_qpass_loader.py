"""Tests for src/qpass_loader.py.

Most tests build tiny fake workbooks in a temporary folder that copy the
quirks of each QPASS format, so they run without the real (gitignored) data.
Tests marked "real data" run only if data/raw/QPASS/ exists locally.
"""

import datetime as dt

import pandas as pd
import pytest
from openpyxl import Workbook

from src.qpass_inventory import RAW_QPASS_DIR, find_header_row, parse_qpass_filename
from src.qpass_loader import (
    DROP_INVALID_DATE,
    DROP_NO_DATE,
    DROP_NO_VALUES,
    DROP_SUMMARY,
    FORMAT_A,
    FORMAT_A_SUMMARY_ROW_LABELS,
    FORMAT_B,
    FORMAT_C,
    OUTPUT_COLUMNS,
    UnsupportedQPASSFormatError,
    _parse_date_cell,
    _week_from_sheet_name,
    load_qpass,
)

# RAW_QPASS_DIR is built from the repo location, so these tests find the real
# data no matter which folder pytest is started from.
REAL_FILES = sorted(RAW_QPASS_DIR.glob("QPASS*.xls*")) if RAW_QPASS_DIR.exists() else []

LONG_HEADER = [
    "Date", "Day", "Zone", "Order", "Report Area", "Neighborhood", "Lot / Structure",
    "Location", "Type", "Report Type", "Total\nStall\nCount",
    "Open Spaces 8am", "Open Spaces 10am", "Open Spaces 12pm", "Open Spaces 2pm",
]
MONDAY = dt.datetime(2024, 9, 30)
TUESDAY = dt.datetime(2024, 10, 1)
# A Monday inside academic year 2026-2027, for the Format C fixture.
FALL_2026_MONDAY = dt.datetime(2026, 9, 28)


def _long_row(date, location, open_8am):
    """One Format B/C data row with fixed placeholder values."""
    return [date, date.strftime("%A"), 1, 1, "Main Campus", "Central", "Structure",
            location, "A", "A", 100, open_8am, 40, 30, 20]


def _save(workbook, folder, filename):
    """Save an openpyxl workbook and return its path."""
    path = folder / filename
    workbook.save(path)
    return path


def _make_format_b(folder, filename="QPASS 2024-2025-2 Fall.xlsx"):
    """Format B: a 'QPASS Data Entry (Week 1)' sheet plus an empty Week 0 sheet."""
    workbook = Workbook()
    week1 = workbook.active
    week1.title = "QPASS Data Entry (Week 1)"
    week1.append(["Name(s):"])  # title row above the header, as in real files
    week1.append(LONG_HEADER)
    week1.append(_long_row(MONDAY, "Athena", 50))
    week1.append(_long_row(TUESDAY, "Athena", 45))
    week0 = workbook.create_sheet("QPASS Data Entry (Week 0)")
    week0.append(["Name(s):"])
    week0.append(LONG_HEADER)  # header only, no data: an empty template
    workbook.create_sheet("Monday (1)").append(["Date", "Space Type"])  # a view, not data
    return _save(workbook, folder, filename)


def _make_format_c(
    folder, filename="QPASS 2026-2027-2  Fall.xlsx", date=FALL_2026_MONDAY
):
    """Format C: one 'Data Entry' sheet with Week # and Sub Type, plus a duplicate Day column."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data Entry"
    header = LONG_HEADER[:2] + ["Week #"] + LONG_HEADER[2:9] + ["Sub Type"] + LONG_HEADER[9:] + ["Day"]
    sheet.append(["Name(s):"])
    sheet.append(header)
    row = _long_row(date, "Gilman", 12)
    sheet.append(row[:2] + ["Week 0"] + row[2:9] + ["24/7 Enforced"] + row[9:] + ["helper value"])
    workbook.create_sheet("Directions & Links").append(["Read me"])
    return _save(workbook, folder, filename)


def _make_format_a(folder, filename="QPASS 2024-2025-2 Fall.xlsx"):
    """Format A: a wide weekly sheet with a multi-row header and blank 'same as above' cells.

    Two days (Monday, Tuesday) x four survey times. Rows:
      Athena / Level 1 / A      -> structure level
      (blank) / Level 2 / B     -> still Athena (structure filled down)
      (blank) / (blank) / S     -> still Athena Level 2 (lot filled down)
      (blank) / P002 / B        -> surface lot: location must be P002, not Athena
      Gilman / (blank) / B      -> structure with no levels: must be Gilman, not P002
      (blank) / (blank) / S     -> still Gilman
      SIO (heading, no numbers) -> dropped
      (blank) / Vacancy         -> summary row, dropped
    """
    times = [dt.time(8), dt.time(10), dt.time(12), dt.time(14)]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "W1 Data"
    sheet.append([None, None, None, None, "Quarterly Parking Availability Survey"])
    sheet.append([None] * 4 + ["Monday", None, None, None, "Tuesday"])
    sheet.append([None] * 4 + [MONDAY, None, None, None, TUESDAY])
    sheet.append(["Structure", "Lot", "Type", "Total Spaces"] + times + times)
    sheet.append([None] * 4 + ["Available"] * 8)
    sheet.append(["Athena", "Level 1", "A", 40, 30, 20, 10, 5, 31, 21, 11, 6])
    sheet.append([None, "Level 2", "B", 50, 40, 30, 20, 10, 41, 31, 21, 11])
    sheet.append([None, None, "S", 60, 50, 40, 30, 20, 51, 41, 31, 21])
    sheet.append([None, "P002", "B", 70, 60, 50, 40, 30, 61, 51, 41, 31])
    sheet.append(["Gilman", None, "B", 80, 70, 60, 50, 40, 71, 61, 51, 41])
    sheet.append([None, None, "S", 90, 80, 70, 60, 50, 81, 71, 61, 51])
    sheet.append(["SIO"])
    sheet.append([None, "Vacancy", None, None, 999, 999, 999, 999, 999, 999, 999, 999])
    workbook.create_sheet("W1 Oversubscription").append(["Structure"])
    return _save(workbook, folder, filename)


# ---------- filename parsing ----------

def test_parse_filename_handles_double_space():
    info = parse_qpass_filename("QPASS 2024-2025-1  Summer.xlsm")
    assert info["academic_year"] == "2024-2025"
    assert info["quarter_number"] == 1
    assert info["quarter_name"] == "Summer"
    assert info["name_matches_number"] is True


def test_parse_filename_flags_number_name_mismatch():
    assert parse_qpass_filename("QPASS 2024-2025-2 Spring.xlsx")["name_matches_number"] is False


@pytest.mark.parametrize("sheet_name, expected", [
    ("QPASS Data Entry (Week 1)", 1),
    ("QPASS Data Entry W2", 2),
    ("W19 W1 Data", 1),          # "W19" means Winter 2019, not week 19
    ("Fall19 Week 2 Data", 2),
    ("F18 Sep17 Data", None),
    ("QPASS Data Entry", None),  # no week number: left blank, not guessed
])
def test_week_from_sheet_name(sheet_name, expected):
    assert _week_from_sheet_name(sheet_name) == expected


# ---------- errors ----------

def test_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_qpass(tmp_path / "QPASS 2024-2025-2 Fall.xlsx")


def test_bad_filename_raises(tmp_path):
    path = _make_format_b(tmp_path, filename="parking data.xlsx")
    with pytest.raises(UnsupportedQPASSFormatError, match="filename"):
        load_qpass(path)


def test_quarter_mismatch_raises(tmp_path):
    path = _make_format_b(tmp_path, filename="QPASS 2024-2025-2 Spring.xlsx")
    with pytest.raises(UnsupportedQPASSFormatError, match="does not match"):
        load_qpass(path)


def test_unknown_layout_raises(tmp_path):
    workbook = Workbook()
    workbook.active.append(["Something else entirely"])
    path = _save(workbook, tmp_path, "QPASS 2024-2025-2 Fall.xlsx")
    with pytest.raises(UnsupportedQPASSFormatError, match="no 'Data Entry' sheet"):
        load_qpass(path)


# ---------- Format B ----------

def test_format_b_loads_rows_and_metadata(tmp_path):
    data = load_qpass(_make_format_b(tmp_path))
    assert list(data.columns) == OUTPUT_COLUMNS
    assert len(data) == 2
    assert (data["format_version"] == FORMAT_B).all()
    assert (data["academic_year"] == "2024-2025").all()
    assert (data["quarter"] == "Fall").all()
    assert (data["source_file"] == "QPASS 2024-2025-2 Fall.xlsx").all()
    assert data["week"].tolist() == [1, 1]
    assert data["open_8am"].tolist() == [50, 45]
    assert data["total_stall_count"].tolist() == [100, 100]  # "Total\nStall\nCount" header matched


def test_format_b_missing_required_column_raises(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QPASS Data Entry (Week 1)"
    sheet.append(["Name(s):"])
    sheet.append([header for header in LONG_HEADER if header != "Open Spaces 8am"])
    sheet.append(_long_row(MONDAY, "Athena", 50)[:-4] + [40, 30, 20])
    path = _save(workbook, tmp_path, "QPASS 2024-2025-2 Fall.xlsx")
    with pytest.raises(UnsupportedQPASSFormatError, match="QPASS 2024-2025-2 Fall.xlsx.*Open Spaces 8am"):
        load_qpass(path)


def test_format_b_counts_undated_leftover_rows(tmp_path):
    """Leftover template rows (no date) are dropped but counted, not silently lost."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QPASS Data Entry (Week 1)"
    sheet.append(["Name(s):"])
    sheet.append(LONG_HEADER)
    sheet.append(_long_row(MONDAY, "Athena", 50))
    sheet.append([None, None, None, None, "East Campus", None, "Lot"])  # leftover row
    data = load_qpass(_save(workbook, tmp_path, "QPASS 2024-2025-2 Fall.xlsx"))
    assert len(data) == 1
    expected = {"QPASS Data Entry (Week 1)": {DROP_NO_DATE: 1}}
    assert data.attrs["dropped_rows"] == expected


def test_format_b_skips_and_records_empty_and_view_sheets(tmp_path):
    skipped = dict(load_qpass(_make_format_b(tmp_path)).attrs["skipped_sheets"])
    assert "empty" in skipped["QPASS Data Entry (Week 0)"]
    assert skipped["Monday (1)"] == "not a data sheet"


# ---------- Format C ----------

def test_format_c_reads_week_and_sub_type(tmp_path):
    data = load_qpass(_make_format_c(tmp_path))
    assert (data["format_version"] == FORMAT_C).all()
    row = data.iloc[0]
    assert row["week"] == 0
    assert row["sub_type"] == "24/7 Enforced"
    assert row["day"] == "Monday"  # first "Day" column, not the helper duplicate
    assert row["source_file"] == "QPASS 2026-2027-2  Fall.xlsx"


def test_sub_type_is_empty_for_formats_without_it(tmp_path):
    data = load_qpass(_make_format_b(tmp_path))
    assert data["sub_type"].isna().all()


def test_format_c_fixture_dates_match_its_academic_year(tmp_path, caplog):
    data = load_qpass(_make_format_c(tmp_path))
    assert data["date"].dt.date.tolist() == [FALL_2026_MONDAY.date()]
    assert "outside academic year" not in caplog.text


def test_dates_outside_academic_year_log_a_warning(tmp_path, caplog):
    """A 2024 date in a file named 2026-2027 is kept but flagged."""
    path = _make_format_c(tmp_path, date=MONDAY)
    data = load_qpass(path)
    assert len(data) == 1
    assert "1 rows have dates outside academic year 2026-2027" in caplog.text


# ---------- Formats B/C: review regressions ----------

def _make_long_sheet(folder, rows, header=LONG_HEADER):
    """Build a Format B workbook with one week sheet holding the given rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QPASS Data Entry (Week 1)"
    sheet.append(["Name(s):"])
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    return _save(workbook, folder, "QPASS 2024-2025-2 Fall.xlsx")


def test_format_b_keeps_text_counts_and_reports_them(tmp_path):
    row = _long_row(MONDAY, "Athena", "FULL")
    data = load_qpass(_make_long_sheet(tmp_path, [row]))
    assert data["open_8am"].tolist() == ["FULL"]
    assert data.attrs["non_numeric_values"] == {"open_8am": 1}


def test_format_b_duplicate_count_column_raises(tmp_path):
    """Regression: a second 'Open Spaces 8am' column must not be silently ignored."""
    header = LONG_HEADER + ["Open Spaces 8am"]
    row = _long_row(MONDAY, "Athena", 50) + [49]
    path = _make_long_sheet(tmp_path, [row], header=header)
    message = r"more than one 'Open Spaces 8am' column \(L, P\)"
    with pytest.raises(UnsupportedQPASSFormatError, match=message):
        load_qpass(path)


def test_excel_serial_date_is_converted_not_read_as_1970(tmp_path):
    """Regression: 45565 is Excel's number for 2024-09-30, not 1970-01-01."""
    row = _long_row(MONDAY, "Athena", 50)
    row[0] = 45565
    data = load_qpass(_make_long_sheet(tmp_path, [row]))
    assert data["date"].dt.date.tolist() == [MONDAY.date()]


def test_invalid_date_is_dropped_and_reported(tmp_path, caplog):
    rows = [_long_row(MONDAY, "Athena", 50), _long_row(MONDAY, "Gilman", 12)]
    rows[1][0] = "TBD"
    data = load_qpass(_make_long_sheet(tmp_path, rows))
    assert data["location"].tolist() == ["Athena"]
    expected = {"QPASS Data Entry (Week 1)": {DROP_INVALID_DATE: 1}}
    assert data.attrs["dropped_rows"] == expected
    warnings = [r.getMessage() for r in caplog.records if r.levelname == "WARNING"]
    assert warnings == [
        "QPASS 2024-2025-2 Fall.xlsx: dropped 1 rows in sheet "
        "'QPASS Data Entry (Week 1)': invalid date"
    ]


@pytest.mark.parametrize("value, expected", [
    (45565, pd.Timestamp("2024-09-30")),
    (45565.75, pd.Timestamp("2024-09-30")),   # a time of day is dropped
    (MONDAY, pd.Timestamp("2024-09-30")),
    ("2024-09-30", pd.Timestamp("2024-09-30")),
    (0, pd.NaT),                              # not a valid Excel date
    (-3, pd.NaT),
    (1e20, pd.NaT),                           # far too large to be a date
    ("Week 0", pd.NaT),
    (None, pd.NaT),
])
def test_parse_date_cell(value, expected):
    result = _parse_date_cell(value)
    if pd.isna(expected):
        assert pd.isna(result)
    else:
        assert result == expected


# ---------- Format A ----------

def test_format_a_reshapes_wide_to_long(tmp_path):
    data = load_qpass(_make_format_a(tmp_path))
    assert (data["format_version"] == FORMAT_A).all()
    # 6 observation rows x 2 days; the heading row and summary row are dropped
    assert len(data) == 12
    assert sorted(data["date"].dt.date.unique()) == [MONDAY.date(), TUESDAY.date()]


def test_format_a_values_land_in_right_date_and_time(tmp_path):
    data = load_qpass(_make_format_a(tmp_path))
    athena = data[data["level"] == "Level 1"]
    athena = athena.set_index(athena["date"].dt.date)
    assert athena.loc[MONDAY.date(), ["open_8am", "open_10am", "open_12pm", "open_2pm"]].tolist() == [30, 20, 10, 5]
    assert athena.loc[TUESDAY.date(), ["open_8am", "open_10am", "open_12pm", "open_2pm"]].tolist() == [31, 21, 11, 6]
    assert athena.loc[TUESDAY.date(), "day"] == "Tuesday"


def test_format_a_fills_structure_only_for_levels(tmp_path):
    monday = load_qpass(_make_format_a(tmp_path))
    monday = monday[monday["date"].dt.date == MONDAY.date()]
    athena_s = monday[monday["total_stall_count"] == 60].iloc[0]  # the Athena "S" row
    assert athena_s["location"] == "Athena"   # blank structure + blank lot -> filled down
    assert athena_s["level"] == "Level 2"
    assert pd.isna(athena_s["raw_lot"])  # raw cell kept blank
    surface = monday[monday["location"] == "P002"]
    assert len(surface) == 1  # P002 is its own lot, not part of Athena
    assert surface["lot_or_structure"].iloc[0] == "Lot"
    assert pd.isna(surface["level"].iloc[0])


def test_format_a_structure_after_surface_lot_keeps_its_name(tmp_path):
    """Regression test: a blank Lot must not inherit the surface lot above it (was P010 bug)."""
    data = load_qpass(_make_format_a(tmp_path))
    monday = data[data["date"].dt.date == MONDAY.date()]
    gilman = monday[monday["total_stall_count"].isin([80, 90])]
    assert gilman["location"].tolist() == ["Gilman", "Gilman"]
    assert gilman["lot_or_structure"].tolist() == ["Structure", "Structure"]
    assert gilman["level"].isna().all()
    assert (monday["location"] == "P002").sum() == 1


def test_format_a_counts_dropped_rows(tmp_path):
    """The 'SIO' heading and the 'Vacancy' row are dropped and counted by reason."""
    data = load_qpass(_make_format_a(tmp_path))
    expected = {"W1 Data": {DROP_NO_VALUES: 1, DROP_SUMMARY: 1}}
    assert data.attrs["dropped_rows"] == expected


def test_format_a_missing_required_label_raises(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "W1 Data"
    sheet.append([None, None, None, MONDAY])
    sheet.append(["Structure", "Lot", "Type", dt.time(8)])  # no "Total Spaces"
    sheet.append(["Athena", "Level 1", "A", 30])
    path = _save(workbook, tmp_path, "QPASS 2024-2025-2 Fall.xlsx")
    with pytest.raises(UnsupportedQPASSFormatError, match="Total Spaces"):
        load_qpass(path)


# The real 2018-2019 files differ in where the label row sits and in column
# order. Each variant: (info columns before the counts, Type as last column?,
# number of title rows above the day row, Excel row of the label row).
FORMAT_A_VARIANTS = {
    "fall_2018": (["Structure", "Lot", "Type", "Total Spaces"], False, 2, 5),
    "summer_2018_no_title": (["Structure", "Lot", "Type", "Total Spaces"], False, 0, 3),
    "2019_2020_total_before_type": (["Structure", "Lot", "Total Spaces", "Type"], False, 2, 5),
    "spring_2019_type_last": (["Structure", "Lot", "Total Spaces"], True, 2, 5),
}


def _make_format_a_variant(folder, info_columns, type_last, title_rows):
    """Build a one-day Format A sheet laid out like one of the real variants."""
    times = [dt.time(8), dt.time(10), dt.time(12), dt.time(14)]
    padding = [None] * len(info_columns)
    tail_label = ["Type"] if type_last else []
    values = {"Structure": "Athena", "Lot": "Level 1", "Type": "A", "Total Spaces": 40}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "W1 Data"
    for _ in range(title_rows):
        sheet.append(padding + ["Quarterly Parking Availability Survey"])
    sheet.append(padding + ["Monday"])
    sheet.append(padding + [MONDAY])
    sheet.append(info_columns + times + tail_label)
    sheet.append(padding + ["Available"] * len(times))
    sheet.append([values[name] for name in info_columns] + [30, 20, 10, 5] + (["A"] if type_last else []))
    return _save(workbook, folder, "QPASS 2024-2025-2 Fall.xlsx")


@pytest.mark.parametrize("variant", FORMAT_A_VARIANTS, ids=str)
def test_format_a_layout_variants_load_the_same_values(tmp_path, variant):
    info_columns, type_last, title_rows, label_excel_row = FORMAT_A_VARIANTS[variant]
    path = _make_format_a_variant(tmp_path, info_columns, type_last, title_rows)
    row = load_qpass(path).iloc[0]
    assert row["location"] == "Athena"
    assert row["level"] == "Level 1"
    assert row["type"] == "A"
    assert row["total_stall_count"] == 40
    assert [row["open_8am"], row["open_10am"], row["open_12pm"], row["open_2pm"]] == [30, 20, 10, 5]
    # the inventory script should point at the same label row
    sheet = pd.read_excel(path, sheet_name="W1 Data", header=None)
    assert find_header_row(sheet) + 1 == label_excel_row


def test_format_a_skips_view_sheets(tmp_path):
    skipped = dict(load_qpass(_make_format_a(tmp_path)).attrs["skipped_sheets"])
    assert skipped == {"W1 Oversubscription": "not a data sheet"}


def _make_format_a_one_day(folder, rows, times=None):
    """Build a one-day Format A sheet (Monday) with the given body rows.

    Each row is [Structure, Lot, Type, Total Spaces, 8am, 10am, 12pm, 2pm].
    """
    times = times or [dt.time(8), dt.time(10), dt.time(12), dt.time(14)]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "W1 Data"
    sheet.append([None] * 4 + ["Monday"])
    sheet.append([None] * 4 + [MONDAY])
    sheet.append(["Structure", "Lot", "Type", "Total Spaces"] + times)
    for row in rows:
        sheet.append(row)
    return _save(workbook, folder, "QPASS 2024-2025-2 Fall.xlsx")


def test_format_a_summary_sections_never_become_locations(tmp_path):
    """Regression: summary labels and the rows under them must not become locations.

    Copies two real layouts: Summer 2018 puts the space type in Structure and
    "Vacancy" in Lot; the 2019 files end with "West Campus" / "SIO" headings
    followed by subtotal rows whose Lot cell is a space type ("A", "SIO B").
    """
    rows = [
        ["Athena", "Level 1", "A", 40, 30, 20, 10, 5],
        [None, "P002", "B", 70, 60, 50, 40, 30],
        ["A", "Vacancy", None, None, 989, 297, 223, 234],
        [None, "Physical Capacity", None, 2217, 2217, 2217, 2217, 2217],
        [None, "Occupancy", None, None, 0.55, 0.87, 0.90, 0.89],
        ["West Campus"],
        [None, "A", None, 1593, 993, 372, 362, 403],
        [None, "Short Term A", None, 56, 48, 23, 16, 18],
        ["SIO"],
        [None, "SIO B", None, 295, 211, 93, 50, 48],
        [None, "Aquarium", None, 258, 240, 130, 23, 48],
    ]
    data = load_qpass(_make_format_a_one_day(tmp_path, rows))
    assert data["location"].tolist() == ["Athena", "P002"]
    expected = {"W1 Data": {DROP_SUMMARY: 7, DROP_NO_VALUES: 2}}
    assert data.attrs["dropped_rows"] == expected


def test_format_a_summary_label_is_matched_without_case(tmp_path):
    rows = [
        [None, "P002", "B", 70, 60, 50, 40, 30],
        [None, "VACANCY ", None, None, 989, 297, 223, 234],
    ]
    data = load_qpass(_make_format_a_one_day(tmp_path, rows))
    assert data["location"].tolist() == ["P002"]


def test_format_a_structure_with_blank_counts_is_not_a_heading(tmp_path):
    """A structure row with a type but no counts must not start a summary section."""
    rows = [
        ["Hopkins", None, "B", None, None, None, None, None],
        [None, "Level 2", "A", 50, 40, 30, 20, 10],
    ]
    data = load_qpass(_make_format_a_one_day(tmp_path, rows))
    assert data["location"].tolist() == ["Hopkins"]
    assert data["level"].tolist() == ["Level 2"]


def test_format_a_keeps_text_counts_and_reports_them(tmp_path):
    """Regression: text counts ("FULL", "n/a") are kept as written, like Formats B/C."""
    rows = [
        ["Hopkins", None, "B", None, "FULL", "FULL", "FULL", "FULL"],
        [None, "P002", "B", 70, "n/a", 50, 40, 30],
    ]
    data = load_qpass(_make_format_a_one_day(tmp_path, rows))
    assert data["location"].tolist() == ["Hopkins", "P002"]
    assert data["open_8am"].tolist() == ["FULL", "n/a"]
    assert data.attrs["dropped_rows"] == {}
    assert data.attrs["non_numeric_values"] == {
        "open_8am": 2, "open_10am": 1, "open_12pm": 1, "open_2pm": 1,
    }


def test_format_a_duplicate_date_and_time_columns_raise(tmp_path):
    """Regression: two Monday 8am columns would let one value overwrite the other."""
    times = [dt.time(8), dt.time(10), dt.time(12), dt.time(14), dt.time(8)]
    rows = [["Athena", "Level 1", "A", 40, 30, 20, 10, 5, 29]]
    path = _make_format_a_one_day(tmp_path, rows, times=times)
    message = r"columns E, I .*2024-09-30 open_8am"
    with pytest.raises(UnsupportedQPASSFormatError, match=message):
        load_qpass(path)


# ---------- real data (runs only if data/raw/QPASS exists) ----------

@pytest.mark.skipif(not REAL_FILES, reason="real QPASS workbooks not downloaded")
@pytest.mark.parametrize("path", REAL_FILES, ids=lambda p: p.name)
def test_real_workbook_loads(path):
    data = load_qpass(path)
    info = parse_qpass_filename(path.name)
    assert list(data.columns) == OUTPUT_COLUMNS
    assert len(data) > 0
    assert (data["academic_year"] == info["academic_year"]).all()
    assert (data["quarter"] == info["quarter_name"]).all()
    assert (data["source_file"] == path.name).all()
    start_year, end_year = (int(year) for year in info["academic_year"].split("-"))
    assert data["date"].dt.year.between(start_year, end_year).all()


# Values that are space types or summary labels, never parking locations.
# They appeared as locations before the summary-section fix (2019 files).
NOT_LOCATIONS = FORMAT_A_SUMMARY_ROW_LABELS | {
    "a", "b", "s", "v", "d", "sio b", "sio s",
}


@pytest.mark.skipif(not REAL_FILES, reason="real QPASS workbooks not downloaded")
@pytest.mark.parametrize("path", REAL_FILES, ids=lambda p: p.name)
def test_real_workbook_has_no_summary_rows_as_locations(path):
    data = load_qpass(path)
    locations = data["location"].astype(str).str.strip().str.casefold()
    assert not locations.isin(NOT_LOCATIONS).any()
    assert data["date"].dt.year.gt(1970).all()
