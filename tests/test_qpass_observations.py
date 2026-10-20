"""Regression tests for task #2's shared-loader observation contract."""

import datetime as dt

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from src.qpass_loader import (
    FORMAT_A, FORMAT_B, FORMAT_C, OUTPUT_COLUMNS, load_all_qpass, load_qpass,
)
from src.qpass_observations import (
    TIME_COLUMNS, reshape_loaded_observations, validate_observations,
)


def table(**changes):
    """Construct one loader-shaped row with distinct category and subtype."""
    row = dict(
        source_file="sample.xlsx", sheet_name="Data Entry", source_row=3,
        format_version=FORMAT_C, academic_year="2026-2027", quarter="Summer",
        date=pd.Timestamp("2026-06-29"), location=" Athena ", type=" Visitor ",
        sub_type="Patient", report_type="Visitor", total_stall_count=10,
        open_8am=3, open_10am=2, open_12pm=1, open_2pm=0,
    )
    row.update(changes)
    return pd.DataFrame([row]).reindex(columns=OUTPUT_COLUMNS)


def audit(wide):
    return validate_observations(reshape_loaded_observations(wide))


def test_four_times_preserve_values_names_metadata_and_input():
    wide = table()
    wide.attrs = {"dropped_rows": {"sample.xlsx": {"no date": 1}}}
    before = wide.copy(deep=True)
    out, summary = audit(wide)
    pd.testing.assert_frame_equal(wide, before)
    assert len(out) == 4
    assert out.observation_datetime.dt.hour.tolist() == [8, 10, 12, 14]
    assert out.open_spaces.tolist() == [3, 2, 1, 0]
    assert out.raw_open_spaces.tolist() == [3, 2, 1, 0]
    assert out["type"].tolist() == [" Visitor "] * 4
    assert out.sub_type.tolist() == ["Patient"] * 4
    assert out.location.tolist() == [" Athena "] * 4
    assert set(OUTPUT_COLUMNS) - set(TIME_COLUMNS) <= set(out.columns)
    assert out.source_row.tolist() == [3] * 4
    assert out.attrs == wide.attrs
    assert out.quality_pass.all()
    assert summary.failed.sum() == 0


@pytest.mark.parametrize("format_version", [FORMAT_A, FORMAT_B, FORMAT_C])
def test_optional_absence_is_not_missing_category(format_version):
    out, _ = audit(table(format_version=format_version, report_type=None,
                         sub_type=None))
    assert out.quality_pass.all()
    expected = "not_in_format" if format_version == FORMAT_A else "blank_or_unavailable"
    assert out.report_type_status.eq(expected).all()
    assert out.sub_type_status.eq(expected).all()
    assert out["type"].eq(" Visitor ").all()


@pytest.mark.parametrize("raw", [None, "", " ", " \t\n "])
@pytest.mark.parametrize("column", ["open_8am", "total_stall_count"])
def test_blank_reading_is_separate_from_text(raw, column):
    out, _ = audit(table(**{column: raw}))
    name = "open_spaces" if column == "open_8am" else column
    assert out.quality_flag.str.contains(f"missing_{name}").any()
    assert not out.quality_flag.str.contains(f"unexpected_text_{name}").any()


@pytest.mark.parametrize("raw", ["FULL", "NA", "n/a", "-", "unknown", "NaN"])
@pytest.mark.parametrize("column", ["open_8am", "total_stall_count"])
def test_text_not_guessed_or_classified_as_blank(raw, column):
    out, _ = audit(table(**{column: raw}))
    name = "open_spaces" if column == "open_8am" else column
    raw_name = "raw_open_spaces" if column == "open_8am" else "raw_total_stall_count"
    assert out.quality_flag.str.contains(f"unexpected_text_{name}").any()
    assert not out.quality_flag.str.contains(f"missing_{name}").any()
    assert out[raw_name].eq(raw).any()


@pytest.mark.parametrize("column,value,flag", [
    ("total_stall_count", 0, "zero_capacity"),
    ("total_stall_count", -1, "negative_total_stall_count"),
    ("total_stall_count", 1.5, "fractional_total_stall_count"),
    ("total_stall_count", np.inf, "nonfinite_total_stall_count"),
    ("open_8am", -1, "negative_open_spaces"),
    ("open_8am", 1.5, "fractional_open_spaces"),
    ("open_8am", np.inf, "nonfinite_open_spaces"),
    ("open_8am", 11, "open_exceeds_capacity"),
])
def test_numeric_quality_checks(column, value, flag):
    out, summary = audit(table(**{column: value}))
    assert out.quality_flag.str.contains(flag).any()
    assert summary.set_index("check").loc[flag, "failed"] > 0
    assert len(out) == 4


@pytest.mark.parametrize("column", ["location", "type", "date"])
def test_required_identity_missing_is_flagged(column):
    out, _ = audit(table(**{column: None}))
    flag = "missing_observation_datetime" if column == "date" else f"missing_{column}"
    assert out.quality_flag.str.contains(flag).all()


def test_undecoded_numeric_date_is_not_epoch_time():
    out, _ = audit(table(date=46202))
    assert out.observation_datetime.isna().all()
    assert out.raw_date.eq(46202).all()


def test_exact_repeats_across_files_preserve_all_rows_and_old_format():
    first = table(format_version=FORMAT_A, report_type=None, sub_type=None,
                  source_row=6)
    second = first.copy()
    second["source_file"] = "copy.xlsx"
    second["source_row"] = 9
    out, summary = audit(pd.concat([first, second], ignore_index=True))
    assert len(out) == 8
    assert out.quality_flag.eq("exact_repeated_record").all()
    assert set(out.source_row) == {6, 9}
    assert not out.quality_flag.str.contains("conflicting_observation").any()
    assert out.observation_group.nunique() == 4
    assert summary.set_index("check").loc["exact_repeated_record", "failed"] == 8


@pytest.mark.parametrize("column", ["total_stall_count", "open_8am"])
def test_conflicting_counts_and_capacity_are_not_exact_repeats(column):
    first = table()
    second = table(**{column: 9, "source_row": 4})
    out, _ = audit(pd.concat([first, second], ignore_index=True))
    expected = 8 if column == "total_stall_count" else 2
    flag = ("ambiguous_allocation_requires_review"
            if column == "total_stall_count" else "conflicting_observation")
    assert out.quality_flag.str.contains(flag).sum() == expected
    if column == "total_stall_count":
        assert not out.quality_flag.str.contains("conflicting_observation").any()
    assert out.quality_flag.str.contains("exact_repeated_record").sum() == 8 - expected
    assert set(out.source_row) == {3, 4}


def test_repeated_subset_of_conflicting_group_gets_both_flags():
    wide = pd.concat([table(), table(source_row=4),
                      table(open_8am=8, source_row=5)], ignore_index=True)
    out, _ = audit(wide)
    morning = out[out.observation_time.eq("08:00")]
    assert morning.quality_flag.str.contains("conflicting_observation").all()
    assert morning.quality_flag.str.contains("exact_repeated_record").tolist() == [
        True, True, False,
    ]
    assert morning.observation_group.nunique() == 1


@pytest.mark.parametrize("column", ["type", "sub_type", "level"])
def test_distinct_allocations_are_not_duplicates(column):
    wide = pd.concat([table(), table(**{column: "Different"})], ignore_index=True)
    out, _ = audit(wide)
    assert out.quality_pass.all()


def test_distinct_invalid_text_and_blank_payloads_are_conflicts():
    out, _ = audit(pd.concat([table(open_8am="FULL"), table(open_8am=" ")]))
    morning = out[out.observation_time.eq("08:00")]
    assert morning.quality_flag.str.contains("conflicting_observation").all()
    assert not morning.quality_flag.str.contains("exact_repeated_record").any()


def test_empty_schema_and_missing_or_duplicate_columns():
    out, summary = audit(table().iloc[:0])
    assert out.empty and "quality_flag" in out
    assert summary[["failed", "passed"]].eq(0).all().all()
    with pytest.raises(ValueError, match="Missing loader columns"):
        reshape_loaded_observations(table().drop(columns=["open_8am"]))
    with pytest.raises(ValueError, match="Duplicate loader columns"):
        reshape_loaded_observations(pd.concat([table(), table()[["type"]]], axis=1))
    with pytest.raises(ValueError, match="Missing observation columns"):
        validate_observations(pd.DataFrame())


def test_summary_overlapping_checks_counts_unique_failed_rows():
    out, summary = audit(table(open_8am=-1, open_10am=11,
                              open_12pm="FULL", open_2pm=1.5))
    counts = summary.set_index("check")
    for flag in ["negative_open_spaces", "open_exceeds_capacity",
                 "unexpected_text_open_spaces", "fractional_open_spaces"]:
        assert counts.loc[flag].tolist() == [1, 3]
    assert counts.loc["any_check"].tolist() == [4, 0]
    assert len(out) == 4


def make_workbook(folder, older=False, weekly=False):
    """Build a raw fixture, not an already-loader-shaped DataFrame."""
    workbook = Workbook()
    sheet = workbook.active
    if older:
        sheet.title = "W1 Data"
        sheet.append([None] * 4 + ["Monday"])
        sheet.append([None] * 4 + [dt.datetime(2018, 10, 1)])
        sheet.append(["Structure", "Lot", "Type", "Total Spaces"] +
                     [dt.time(8), dt.time(10), dt.time(12), dt.time(14)])
        sheet.append(["Athena", "Level 1", "A", 10, 3, "FULL", " ", 0])
        sheet.append([None, "Level 1", "A", 10, 3, "FULL", " ", 0])
        filename = "QPASS 2018-2019-2 Fall.xlsx"
    else:
        sheet.title = "QPASS Data Entry (Week 1)" if weekly else "Data Entry"
        sheet.append(["Title"])
        sheet.append(["Date", "Location", "Type", "Sub Type", "Report Type",
                      "Total Stall Count", "Open Spaces 8am", "Open Spaces 10am",
                      "Open Spaces 12pm", "Open Spaces 2pm"])
        sheet.append([dt.datetime(2026, 6, 29), "Athena", "Visitor", "Patient",
                      "Visitor", 10, 3, "n/a", " ", 0])
        # Formula-like template placeholder: only shared loader decides exclusion.
        sheet.append([dt.datetime(2026, 6, 29), 0, 0, None, None, 0, 0, 0, 0, 0])
        filename = "QPASS 2026-2027-1 Summer.xlsx"
    path = folder / filename
    workbook.save(path)
    return path


@pytest.mark.parametrize("older,weekly", [(True, False), (False, True), (False, False)])
def test_shared_loader_reshape_validate_integration_all_formats(tmp_path, older, weekly):
    loaded = load_qpass(make_workbook(tmp_path, older=older, weekly=weekly))
    before = loaded.copy(deep=True)
    out, summary = audit(loaded)
    pd.testing.assert_frame_equal(loaded, before)
    assert len(out) == 4 * len(loaded)
    assert out.attrs == loaded.attrs
    assert out.quality_flag.str.contains("unexpected_text_open_spaces").any()
    assert out.quality_flag.str.contains("missing_open_spaces").any()
    if older:
        assert out.format_version.eq(FORMAT_A).all()
        assert out.report_type_status.eq("not_in_format").all()
        assert out[out.observation_time.eq("08:00")].quality_flag.eq(
            "exact_repeated_record"
        ).all()
        assert out["type"].eq("A").all()
        assert set(out.source_row) == {4, 5}
    else:
        assert len(loaded) == 1
        assert loaded.attrs["dropped_rows"]
        assert out.sub_type.eq("Patient").all()
    assert summary.set_index("check").loc["any_check", "failed"] > 0


def test_combined_loader_preserves_per_file_metadata(tmp_path):
    make_workbook(tmp_path, older=True)
    make_workbook(tmp_path)
    loaded = load_all_qpass(tmp_path)
    out, _ = audit(loaded)
    assert len(out) == 4 * len(loaded)
    assert set(out.academic_year) == {"2018-2019", "2026-2027"}
    assert out.attrs == loaded.attrs
    assert set(out.source_file) == set(loaded.source_file)


@pytest.mark.parametrize("weekly", [True, False])
@pytest.mark.parametrize("header,column", [
    ("Sub Type", "sub_type"), ("Report Type", "report_type"),
])
def test_optional_headers_absent_versus_blank_per_sheet(tmp_path, weekly, header, column):
    """Header absence survives mixed sheets and combined-file loading."""
    from openpyxl import load_workbook

    path = make_workbook(tmp_path, weekly=weekly)
    book = load_workbook(path)
    sheet = book.active
    position = next(cell.column for cell in sheet[2] if cell.value == header)
    sheet.cell(3, position).value = None
    if weekly:
        absent = book.copy_worksheet(sheet)
        absent.title = "QPASS Data Entry (Week 2)"
        absent.delete_cols(position)
    else:
        sheet.delete_cols(position)
    book.save(path)
    loaded = load_all_qpass(tmp_path)
    out, _ = audit(loaded)
    absent_rows = out if not weekly else out[out.sheet_name.str.endswith("(Week 2)")]
    assert absent_rows[f"{column}_status"].eq("not_in_format").all()
    assert absent_rows[f"{column}_in_format"].eq(False).all()
    if weekly:
        blank_rows = out[out.sheet_name.str.endswith("(Week 1)")]
        assert blank_rows[f"{column}_status"].eq("blank").all()
        assert blank_rows[f"{column}_in_format"].eq(True).all()
    assert not out.quality_flag.str.contains(f"missing_{column}").any()
    assert out.attrs == loaded.attrs


def test_ambiguous_capacity_group_retains_repeats_counts_and_provenance():
    wide = pd.concat([table(), table(source_row=4),
                      table(total_stall_count=20, open_8am=8, source_row=5)],
                     ignore_index=True)
    out, summary = audit(wide)
    assert len(out) == 12
    assert out.quality_flag.str.contains("ambiguous_allocation_requires_review").all()
    assert not out.quality_flag.str.contains("conflicting_observation").any()
    assert out.quality_flag.str.contains("exact_repeated_record").sum() == 8
    assert out.observation_group.nunique() == 4
    assert set(out.source_row) == {3, 4, 5}
    assert set(out.total_stall_count) == {10, 20}
    assert summary.set_index("check").loc[
        "ambiguous_allocation_requires_review", "failed"
    ] == 12


@pytest.mark.parametrize("format_version", [FORMAT_A, FORMAT_B])
def test_other_format_capacity_variants_remain_conflicts(format_version):
    out, _ = audit(pd.concat([
        table(format_version=format_version),
        table(format_version=format_version, total_stall_count=20),
    ], ignore_index=True))
    assert out.quality_flag.str.contains("conflicting_observation").all()
    assert not out.quality_flag.str.contains("ambiguous_allocation_requires_review").any()


@pytest.mark.parametrize("older,weekly", [(True, False), (False, True), (False, False)])
def test_audited_observations_integrate_with_main_label_adapter(tmp_path, older, weekly):
    from src.labels import prepare_label_records

    loaded = load_qpass(make_workbook(tmp_path, older=older, weekly=weekly))
    audited, _ = audit(loaded)
    before = audited.copy(deep=True)
    labels = prepare_label_records(audited)
    pd.testing.assert_frame_equal(audited, before)
    pd.testing.assert_frame_equal(labels[audited.columns], audited)
    for field, source in [("facility_raw", "location"), ("category_raw", "type"),
                          ("subtype_raw", "sub_type"), ("source_sheet", "sheet_name")]:
        pd.testing.assert_series_equal(labels[field], audited[source], check_names=False)
    assert labels.attrs == audited.attrs


def test_source_order_preserved_without_assuming_allocation_identity(tmp_path):
    from openpyxl import load_workbook

    path = make_workbook(tmp_path)
    book = load_workbook(path)
    sheet = book.active
    sheet.cell(2, 11, "Order")
    sheet.cell(3, 11, 7)
    book.save(path)
    loaded = load_qpass(path)
    out, _ = audit(loaded)
    assert out.source_order.eq(7).all()
